"""
文档处理集成测试 - 测试端到端的文档处理流程
"""

import os
import csv
import pytest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from src.application import process_single_file
from src.processor.doc_processor import DocProcessor, process_doc
from tests.utils.test_helpers import create_sample_document


class TestIntegrationDocProcessing:
    """测试端到端文档处理流程"""

    @pytest.fixture
    def sample_doc_path(self, tmpdir):
        """创建测试文档"""
        doc_path = os.path.join(tmpdir, "test_doc.docx")
        return create_sample_document(doc_path, batch_number="1")

    def test_end_to_end_processing(self, sample_doc_path, tmpdir):
        """测试端到端处理流程"""
        # 处理文档
        processor = DocProcessor(sample_doc_path)
        processor.process()

        # 保存处理结果
        output_path = os.path.join(tmpdir, "output.csv")
        processor.save_to_csv(output_path)

        # 验证结果
        assert os.path.exists(output_path)
        assert processor.cars
        assert len(processor.cars) == 3  # 默认3个记录

        # 验证处理后的数据结构
        for car in processor.cars:
            assert "vmodel" in car
            assert "企业名称" in car
            assert "品牌" in car

    @patch("src.processor.doc_processor.verify_batch_consistency")
    def test_process_with_batch_verification(self, mock_verify, sample_doc_path):
        """测试带批次验证的处理流程"""
        # 设置mock返回值
        mock_verify.return_value = {"status": "OK", "details": {}}

        # 处理文档
        processor = DocProcessor(sample_doc_path)
        processor.process()

        # 验证批次验证被调用
        mock_verify.assert_called_once()

    def test_process_doc_function(self, sample_doc_path, tmpdir):
        """测试process_doc便捷函数"""
        # 设置输出路径
        output_path = os.path.join(tmpdir, "output.csv")

        # 使用process_doc函数处理文档
        cars = process_doc(sample_doc_path, output_path)

        # 验证结果
        assert cars
        assert len(cars) == 3
        assert os.path.exists(output_path)


def test_invalid_rows_are_exported_with_source_location(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    output_dir = tmp_path / "output"
    document = Document()
    document.add_paragraph("第1批")
    document.add_paragraph("新能源汽车")
    table = document.add_table(rows=3, cols=4)
    for cell, value in zip(
        table.rows[0].cells, ["序号", "企业名称", "品牌", "车辆型号"]
    ):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["1", "企业A", "品牌A", "MODEL-A"]):
        cell.text = value
    for cell, value in zip(table.rows[2].cells, ["2", "企业B", "品牌B", ""]):
        cell.text = value
    document.save(source)

    records = process_single_file(
        str(source), str(output_dir), verbose=False, display=False
    )

    assert len(records) == 1
    invalid_output = output_dir / "sample_invalid_records.csv"
    assert invalid_output.exists()
    with invalid_output.open(encoding="utf-8-sig", newline="") as handle:
        invalid_rows = list(csv.DictReader(handle))
    assert len(invalid_rows) == 1
    assert invalid_rows[0]["source_file"] == "sample.docx"
    assert invalid_rows[0]["table_id"] == "1"
    assert invalid_rows[0]["row_number"] == "3"
    assert invalid_rows[0]["序号"] == "2"
    assert invalid_rows[0]["企业名称"] == "企业B"
    assert invalid_rows[0]["reason"] == "缺少必要字段: vmodel"
