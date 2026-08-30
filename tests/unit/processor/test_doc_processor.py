"""Focused tests for the core document processor."""

import csv
from pathlib import Path
from typing import Optional
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from src.processor.doc_processor import (
    DocProcessor,
    DocumentError,
    ProcessingError,
    process_doc,
)


def create_vehicle_document(
    path: Path,
    include_invalid_row: bool = False,
    declared_count: Optional[int] = None,
) -> None:
    document = Document()
    document.add_paragraph("第1批")
    document.add_paragraph("新能源汽车")
    document.add_paragraph("（一）乘用车")
    if declared_count is not None:
        document.add_paragraph(f"本批次共计{declared_count}款车型")
    row_count = 3 if include_invalid_row else 2
    table = document.add_table(rows=row_count, cols=4)
    for cell, value in zip(table.rows[0].cells, ["序号", "企业名称", "品牌", "型号"]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["1", "企业A", "品牌A", "MODEL-A"]):
        cell.text = value
    if include_invalid_row:
        for cell, value in zip(table.rows[2].cells, ["2", "企业B", "品牌B", ""]):
            cell.text = value
    document.save(path)


def test_init_applies_verification_config(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)
    processor = DocProcessor(
        str(source),
        verbose=False,
        config={
            "performance": {"chunk_size": 25},
            "document": {"skip_count_check": True, "skip_verification": True},
        },
    )
    assert processor._chunk_size == 25
    assert processor._skip_count_check is True
    assert processor._skip_verification is True


def test_invalid_document_raises_document_error(tmp_path: Path) -> None:
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"not a zip package")
    with pytest.raises(DocumentError, match="无法加载文档"):
        DocProcessor(str(source))


def test_streaming_and_standard_modes_produce_identical_results(
    tmp_path: Path,
) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source, include_invalid_row=True, declared_count=2)
    standard = DocProcessor(
        str(source),
        verbose=False,
        config={"performance": {"streaming_xml": False}},
    )
    streaming = DocProcessor(
        str(source),
        verbose=False,
        config={"performance": {"streaming_xml": True}},
    )

    assert streaming.doc is None
    assert streaming.process() == standard.process()
    assert streaming.candidate_record_count == standard.candidate_record_count == 2
    assert streaming.invalid_record_count == standard.invalid_record_count == 1
    assert streaming.declared_count == standard.declared_count == 2
    assert streaming.consistency_result == standard.consistency_result


def test_auto_streaming_threshold_can_select_streaming(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)

    processor = DocProcessor(
        str(source),
        config={
            "performance": {
                "streaming_xml": "auto",
                "streaming_threshold_mb": 0,
            }
        },
    )

    assert processor._streaming_xml is True
    assert processor.doc is None


def test_invalid_streaming_config_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)

    with pytest.raises(DocumentError, match="streaming_xml"):
        DocProcessor(
            str(source),
            config={"performance": {"streaming_xml": "sometimes"}},
        )


def test_process_tracks_candidates_and_invalid_rows(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source, include_invalid_row=True)
    processor = DocProcessor(str(source), verbose=False)

    records = processor.process()

    assert len(records) == 1
    assert processor.candidate_record_count == 2
    assert processor.invalid_record_count == 1
    assert processor.consistency_result["status"] == "internal_mismatch"
    assert processor.consistency_result["difference"] == 1


def test_skip_verification_is_honored(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)
    processor = DocProcessor(
        str(source),
        verbose=False,
        config={"document": {"skip_verification": True}},
    )

    processor.process()

    assert processor.consistency_result["status"] == "skipped"
    assert processor.declared_count is None


def test_save_to_csv_is_complete_and_quoted(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)
    processor = DocProcessor(str(source), verbose=False)
    processor.cars = [
        {"vmodel": "MODEL,1", "category": "新能源"},
        {"vmodel": "MODEL-2", "late_field": "late value"},
    ]
    output = tmp_path / "nested" / "result.csv"

    processor.save_to_csv(str(output))

    with output.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["vmodel"] == "MODEL,1"
    assert rows[1]["late_field"] == "late value"
    assert not list(output.parent.glob("*.tmp"))


@patch("src.processor.doc_processor.DocProcessor")
def test_process_doc_saves_result(mock_processor_class: MagicMock) -> None:
    processor = mock_processor_class.return_value
    records = [{"vmodel": "MODEL-1"}]
    processor.process.return_value = records

    assert process_doc("sample.docx", "result.csv") == records
    processor.save_to_csv.assert_called_once_with("result.csv")


@patch("src.processor.doc_processor.DocProcessor")
def test_process_doc_propagates_document_errors(
    mock_processor_class: MagicMock,
) -> None:
    mock_processor_class.side_effect = DocumentError("broken")
    with pytest.raises(DocumentError, match="broken"):
        process_doc("invalid.docx")


@patch("src.processor.doc_processor.psutil")
def test_get_memory_usage(mock_psutil: MagicMock, tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)
    mock_psutil.Process.return_value.memory_info.return_value.rss = 50 * 1024 * 1024
    processor = DocProcessor(str(source))
    assert processor.get_memory_usage() == "50.0MB"


def test_processing_error_is_not_double_wrapped(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    create_vehicle_document(source)
    processor = DocProcessor(str(source), verbose=False)
    processor.table_extractor.extract_car_info = MagicMock(
        side_effect=RuntimeError("table failure")
    )
    with pytest.raises(ProcessingError, match="元素错误"):
        processor.process()
